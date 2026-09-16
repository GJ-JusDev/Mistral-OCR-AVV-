import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    return heading

doc = Document()

# Add a title
title = doc.add_heading('Chapter 3\nDISCUSSION OF FINDINGS', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    "This chapter presents and discusses the requirements for the DocuValidate (Mistral OCR) application process. "
    "It covers the essential information, recommended model for the system, proposed features, and desired usability level. "
    "Additionally, this chapter will delve into the intricacies of system design and physical design, elucidating essential "
    "factors and considerations necessary for the effective process of the proposed system."
)

doc.add_heading('Information, Hardware, and Software Requirements of the Proposed System', level=2)
doc.add_paragraph(
    "The DocuValidate system replaces manual document verification with an Automated OCR and Validation System. "
    "The following are the key information requirements needed to design and implement this system effectively:"
)

doc.add_paragraph("1. User & Role Information. Vital for managing access control (Admin vs Teacher). This includes auth logs, roles, and teacher invite codes.")
doc.add_paragraph("2. Student Information. The system tracks student details like LRN, name, birth date, grade level, and school ID to link them with uploaded documents.")
doc.add_paragraph("3. Document OCR Information. Captures image URLs, original filenames, OCR model used, confidence scores, and extracted JSON fields.")
doc.add_paragraph("4. Validation & Review Information. Stores automated validation results (rule name, passed status, actual/expected values) and manual teacher review logs.")

doc.add_heading('Software Requirements', level=3)
doc.add_paragraph("The following are the software requirements for the DocuValidate System:")
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Item/s'
hdr_cells[1].text = 'Minimum Requirement'
hdr_cells[2].text = 'Recommended'

reqs = [
    ('Operating System', 'Windows 10 / macOS', 'Windows 11 / Latest macOS'),
    ('Local Environment', 'Node.js LTS (v18)', 'Node.js LTS (v20+)'),
    ('Code Editor', 'Visual Studio Code', 'Visual Studio Code (Latest)'),
    ('Web Browsers', 'Chrome, Firefox', 'Latest Chrome or Firefox'),
    ('Web Framework', 'Next.js 14', 'Next.js 14 (App Router)'),
    ('Database', 'PostgreSQL (Supabase)', 'Supabase Managed PostgreSQL'),
    ('External APIs', 'Mistral AI API, Gmail SMTP', 'Mistral AI API, Gmail SMTP'),
]
for item, min_req, rec in reqs:
    row_cells = table.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = min_req
    row_cells[2].text = rec

doc.add_heading('Hardware Requirements', level=3)
table2 = doc.add_table(rows=1, cols=3)
table2.style = 'Table Grid'
hdr_cells2 = table2.rows[0].cells
hdr_cells2[0].text = 'Item/s'
hdr_cells2[1].text = 'Minimum Requirement'
hdr_cells2[2].text = 'Recommended'

h_reqs = [
    ('Processor', 'Intel Core i3 / AMD Ryzen 3', 'Intel Core i5 / Apple M1 or higher'),
    ('Memory (RAM)', '8 GB RAM', '16 GB RAM or higher'),
    ('Storage', '256 GB SSD', '512 GB SSD or higher'),
    ('Internet Connection', '10 Mbps', '50 Mbps or higher for fast image uploads'),
]
for item, min_req, rec in h_reqs:
    row_cells = table2.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = min_req
    row_cells[2].text = rec

doc.add_heading('Flow of Data of the Proposed System', level=2)
doc.add_paragraph(
    "Manual Process: Traditionally, teachers manually receive student documents, verify the content against the student directory, "
    "and record the results in physical logbooks or spreadsheets. This is prone to human error and is time-consuming.\n\n"
    "Proposed Process: In the DocuValidate system, teachers use their device's camera or upload an image of the document. "
    "The system calls the Mistral OCR API to extract text and automatically parses it into structured fields. "
    "The validation engine then compares these fields against existing database records. The results are instantly stored and presented to the teacher for review."
)

doc.add_heading('Features of the proposed system', level=2)
doc.add_paragraph("The DocuValidate system is designed to streamline administrative validation tasks:")
doc.add_paragraph("- Dashboard: Displays system statistics and authentication logs for Admins.")
doc.add_paragraph("- Document Upload & OCR: Integrates hardware camera and TrOCR/Mistral AI to process images directly from the browser.")
doc.add_paragraph("- Review Queue: A structured interface where teachers review AI-extracted fields, see validation discrepancies, and approve or reject documents.")
doc.add_paragraph("- Settings & Role Management: Secure Admin panel to generate teacher access codes and manage school data.")

doc.add_heading('Extent of Usability of the Proposed System', level=2)
doc.add_paragraph(
    "To evaluate the overall usability, a System Usability Scale (SUS) questionnaire was administered to selected users (Admins and Teachers). "
    "The SUS consists of ten statements on a five-point Likert scale. Based on initial testing, the system received highly positive feedback, "
    "particularly for its automated OCR extraction, which drastically reduced manual data entry time."
)

doc.add_heading('Systems Architecture', level=2)
doc.add_paragraph(
    "The system adopts a modern client-server architectural pattern using Next.js App Router. "
    "It separates presentation (Client Components), server logic (Server Actions and API routes), "
    "and the backend database (Supabase PostgreSQL). Secure communication is ensured via HTTPS."
)

doc.add_heading('Architectural Views', level=3)

doc.add_heading('Use Case Diagram', level=4)
doc.add_paragraph(
    "The Use Case Diagram provides a clear overview of the functional requirements. "
    "It outlines the roles of Admin and Teacher. Both can access the Dashboard, Upload Document, Review Queue, and Student Directory. "
    "However, Admins have extended permissions to view Auth Logs and generate Teacher Invites in the Settings."
)

doc.add_heading('Class Diagram / Data Model', level=4)
doc.add_paragraph(
    "The Class Diagram maps the Supabase database schema. Key entities include 'students', 'documents', "
    "'validation_results', 'review_logs', 'user_roles', and 'teacher_invites'. Documents are linked to students via foreign keys, "
    "while automated validation and manual review logs are tied directly to the documents."
)

doc.add_heading('Sequence Diagrams', level=4)
doc.add_paragraph(
    "Document Upload & OCR Sequence: The user client invokes the `useOcr` hook, sending an image to the Next.js API. "
    "The API acts as a secure proxy to the Mistral AI API. The returned markdown is parsed by the Validation Engine into structured fields and saved to Supabase.\n\n"
    "Teacher Registration Sequence: A new teacher submits an email, password, and access code. A Next.js Server Action validates the code in the DB. "
    "If valid, the user is created via Supabase Auth, triggering a PostgreSQL function that automatically assigns the 'teacher' role."
)

doc.add_heading('Deployment Diagram', level=4)
doc.add_paragraph(
    "The application is deployed across managed cloud environments. Vercel hosts the Next.js frontend and serverless API endpoints. "
    "Supabase provides the managed PostgreSQL database and GoTrue Auth Server. External APIs include Mistral AI (for OCR) and Gmail SMTP (for emailing invites)."
)

doc.add_page_break()

# Chapter 4
title4 = doc.add_heading('Chapter 4\nCONCLUSIONS AND RECOMMENDATIONS', level=1)
title4.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('Conclusion', level=2)
doc.add_paragraph("After developing and reviewing the DocuValidate system, the researchers concluded the following:")
doc.add_paragraph("1. Manual document verification is inefficient. A digital OCR-based system is essential to reduce administrative burden.")
doc.add_paragraph("2. The Next.js and Supabase architecture provided a secure, role-based platform that successfully differentiated Admin and Teacher workflows.")
doc.add_paragraph("3. The integration of Mistral AI OCR, coupled with custom validation logic, effectively replaced manual data entry, streamlining the school's operations.")
doc.add_paragraph("4. The system improved overall efficiency and user satisfaction, with users finding the review queue and automated confidence scoring reliable and easy to use.")

doc.add_heading('Recommendations', level=2)
doc.add_paragraph("Based on the findings and conclusions of this study, the following recommendations are proposed to improve the system:")
doc.add_paragraph("1. Enhance mobile responsiveness across the camera and cropping UI for better handling on lower-end devices.")
doc.add_paragraph("2. Implement regular security audits and fine-tune Supabase Row Level Security (RLS) policies as the user base grows.")
doc.add_paragraph("3. Integrate additional local or fallback OCR models to reduce dependency on external APIs and improve latency.")
doc.add_paragraph("4. Expand reporting capabilities to generate automated weekly or monthly summaries of validated versus rejected documents.")

doc.save('c:/Users/windows-11/Desktop/Mistral Ocr AAV Finale/Chapter3_4_DocuValidate.docx')
print("DOCX generated successfully.")
